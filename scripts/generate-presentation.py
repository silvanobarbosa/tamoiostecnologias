"""Generate editable PowerPoint from the same element registry as Tamoios Slides."""
import argparse
import json
import re
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]

def generate(destination):
    guide = json.loads((ROOT / "src/lib/brand-guide.json").read_text())
    colors = {c["name"]: c["hex"].lstrip("#") for c in guide["colors"]}
    slides = json.loads((ROOT / "src/lib/presentation-layouts.json").read_text())
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Tamoios Tecnologia — Biblioteca de modelos v2"
    prs.core_properties.author = "Tamoios Tecnologia"
    prs.core_properties.subject = "24 modelos editáveis de apresentação"
    unit = lambda px: Inches(px / 144)
    for layout in slides:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string(colors[layout["background"]])
        for e in layout["elements"]:
            x, y, w, h = [unit(e[k]) for k in ("x", "y", "w", "h")]
            if e["type"] == "rect":
                shape = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(colors[e["fill"]])
                shape.line.fill.background()
            elif e["type"] == "image":
                path = ROOT / "src/assets" / e["asset"]
                iw, ih = Image.open(path).size
                ratio = iw / ih
                target = e["w"] / e["h"]
                if e.get("fit", "contain") == "contain":
                    dw = min(e["w"], e["h"] * ratio)
                    dh = dw / ratio
                    s.shapes.add_picture(str(path), unit(e["x"] + (e["w"] - dw) / 2), unit(e["y"] + (e["h"] - dh) / 2), unit(dw), unit(dh))
                else:
                    pic = s.shapes.add_picture(str(path), x, y, w, h)
                    if ratio > target:
                        crop = (1 - target / ratio) / 2
                        pic.crop_left = crop
                        pic.crop_right = crop
                    else:
                        crop = (1 - ratio / target) / 2
                        pic.crop_top = crop
                        pic.crop_bottom = crop
            elif e["type"] == "text":
                box = s.shapes.add_textbox(x, y, w, h)
                tf = box.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.word_wrap = True
                for idx, line in enumerate(e.get("text", "").split("\n")):
                    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                    p.text = line
                    p.font.name = "Fraunces" if e.get("font") == "display" else "Inter"
                    p.font.size = Pt(e.get("size", 38) / 2)
                    p.font.bold = e.get("bold", False)
                    p.font.color.rgb = RGBColor.from_string(colors[e.get("color", "forest")])
                    p.line_spacing = 1.1 if e.get("font") == "display" else 1.32
                    p.space_before = p.space_after = Pt(0)
        s.notes_slide.notes_text_frame.text = layout["notes"]
    destination.parent.mkdir(parents=True, exist_ok=True)
    prs.save(destination)
    print(f"Created {len(slides)} editable slides: {destination}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=ROOT / "public/docs/identidade/tamoios-template-comercial-v2.pptx")
    generate(parser.parse_args().output)